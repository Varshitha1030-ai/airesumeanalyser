"""Extract plain text from uploaded resumes (PDF or DOCX)."""
import io

import pdfplumber
from docx import Document


def extract_pdf(file) -> str:
    """Return all text from a PDF. `file` can be a path or a file-like object."""
    pages = []
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n".join(pages).strip()


def extract_docx(file) -> str:
    """Return all text from a DOCX (paragraphs and table cells)."""
    doc = Document(file)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts).strip()


def extract_text(uploaded_file) -> str:
    """Pick the right extractor from the file name.

    Works with a Streamlit UploadedFile (has .name and .getvalue()).
    """
    name = uploaded_file.name.lower()
    data = io.BytesIO(uploaded_file.getvalue())
    if name.endswith(".pdf"):
        return extract_pdf(data)
    if name.endswith(".docx"):
        return extract_docx(data)
    raise ValueError("Unsupported file type. Please upload a PDF or DOCX.")
